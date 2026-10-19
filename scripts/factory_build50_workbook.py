#!/usr/bin/env python3
"""One real 50-stage technical workbook producer, not a 40-family inventor."""
import argparse
import ast
import copy
from fractions import Fraction
import hashlib
import importlib.util
import io
import json
import math
import os
from pathlib import Path
import shutil
import subprocess
import sys
import unittest
import zipfile

spec=importlib.util.spec_from_file_location('workbook_build50',Path(__file__).with_name('factory_build_50.py'))
protocol=importlib.util.module_from_spec(spec);spec.loader.exec_module(protocol)
PRODUCT='capacity_planning_workbook'
# Formula implementation and reference functions are separate; LibreOffice is a third engine.
UNITS=[
 ('Work capacity','a*b*c',[4,6,8],['People','Units per person per hour','Hours'],['people','units/person/hour','hours'],'units',lambda a,b,c:a*b*c),
 ('Material demand','a*b',[120,2.5],['Planned units','Material per unit'],['units','kg/unit'],'kg',lambda a,b:a*b),
 ('Replicated storage','a*b*c/1024',[200,8,3],['Files','Size per file','Replicas'],['files','MiB/file','copies'],'GiB',lambda a,b,c:a*b*c/Fraction(1024)),
 ('Backup duration','a*1024/b',[12,24],['Data volume','Transfer rate'],['GiB','MiB/second'],'seconds',lambda a,b:a*Fraction(1024)/b),
 ('Batch throughput','a*b/c',[8,20,30],['Batches','Units per batch','Elapsed minutes'],['batches','units/batch','minutes'],'units/minute',lambda a,b,c:a*b/c),
 ('Capacity load','a/b*100',[28,40],['Required work hours','Available hours'],['hours','hours'],'percent',lambda a,b:Fraction(100)*a/b),
 ('Defect frequency','a/b*100',[3,120],['Recorded defects','Processed jobs'],['defects','jobs'],'defects per 100 jobs',lambda a,b:Fraction(100)*a/b),
 ('Repair duration','a*b/c',[12,20,3],['Defects','Minutes per repair','Parallel workers'],['defects','minutes/defect','workers'],'minutes',lambda a,b,c:a*b/c),
 ('Queue delay','a/b',[60,4],['Pending jobs','Processing rate'],['jobs','jobs/minute'],'minutes',lambda a,b:a/b),
 ('Energy estimate','a*b/1000',[600,5],['Power','Operating hours'],['watts','hours'],'kWh',lambda a,b:a*b/Fraction(1000)),
 ('Water volume','a*b',[12,3],['Flow rate','Operating hours'],['litres/hour','hours'],'litres',lambda a,b:a*b),
 ('Inventory coverage','a/b',[240,12],['Available inventory','Daily consumption'],['units','units/day'],'days',lambda a,b:a/b),
]
ACTIONS=('Specify finite input domains and meaningful units for ',
         'Construct an executable worksheet with transparent formulas for ',
         'Independently test normal, zero and scaled reference cases for ',
         'Revise the worksheet with protected formulas, guidance and validation for ')

def valid(unit,values):
    if len(values)!=len(unit[2]) or any(type(x) not in (int,float) or not math.isfinite(x) or not 0<=x<=1000000 for x in values):
        raise ValueError('INVALID_QUANTITY')
    names=dict(zip('abc',values))
    tree=ast.parse(unit[1],mode='eval')
    for node in ast.walk(tree):
        if not isinstance(node,(ast.Expression,ast.BinOp,ast.Mult,ast.Div,ast.Name,ast.Load,ast.Constant)):
            raise ValueError('UNSAFE_FORMULA')
    for node in ast.walk(tree):
        if isinstance(node,ast.BinOp) and isinstance(node.op,ast.Div) and isinstance(node.right,ast.Name):
            if names[node.right.id]<1e-9:raise ValueError('ZERO_OR_UNSUPPORTED_DENOMINATOR')
    return names,tree

def reference(unit,values):
    valid(unit,values)
    return float(unit[6](*[Fraction(str(x)) for x in values]))

def checks(unit):
    count=0
    for scale in (0.5,1,2,10):
        values=[x*scale for x in unit[2]];names,tree=valid(unit,values)
        measured=eval(compile(tree,'<fixed-formula>','eval'),{'__builtins__':{}},names)
        assert math.isclose(measured,reference(unit,values),rel_tol=1e-10,abs_tol=1e-10);count+=1
    values=list(unit[2]);values[0]=0;names,tree=valid(unit,values)
    assert eval(compile(tree,'<fixed-formula>','eval'),{'__builtins__':{}},names)==reference(unit,values);count+=1
    for bad in (True,float('inf'),-1):
        values=list(unit[2]);values[0]=bad
        try:valid(unit,values)
        except ValueError:count+=1
        else:raise ValueError('NEGATIVE_CASE_NOT_REJECTED')
    return count

def build(root,tools):
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Font,PatternFill,Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    root=Path(root)
    profile_data=[{'title':u[0],'formula':u[1],'values':u[2],'labels':u[3],'units':u[4],'result_unit':u[5],'maximum_input':1000000,'minimum_positive_divisor':1e-9} for u in UNITS]
    source=protocol.digest({'profiles':profile_data,'producer_sha256':protocol.file_sha(__file__),
                            'protocol_sha256':protocol.file_sha(Path(protocol.__file__))})
    plan={'schema':protocol.SCHEMA,'product_id':PRODUCT,'producer_id':'openpyxl_workbook_builder',
          'source_snapshot_sha256':source,'work_packages':[
              {'id':'unit_%02d_%d'%(n,action),'instruction':ACTIONS[action]+unit[0]+'.',
               'acceptance':['Required inputs, units and outputs are explicit',
                             'The stage output passes its independent acceptance checks']}
              for n,unit in enumerate(UNITS) for action in range(4)]}
    home=root/'build-50'/PRODUCT;state_path=home/'state.json';target=root/'digital-products'/PRODUCT
    if state_path.exists():
        state=protocol.read(state_path)
        if state.get('status')=='awaiting_owner_acceptance' and protocol.verify_attempt(home,state,source):
            try:
                manifest=protocol.read(target/'manufacturing_manifest.json')
                if (protocol.file_sha(target/(PRODUCT+'-internal-qa.zip'))==manifest['artifact_sha256']
                        and manifest['artifact_sha256']==protocol.file_sha(Path(state['attempt_folder'])/(PRODUCT+'-internal-qa.zip'))):
                    return {'status':'unchanged_verified_build','completed_stages':50,'commercial_release':False}
            except (ValueError,OSError,KeyError):pass
    def result(names,count=2):return {'passed':True,'outputs':names,'successful_checks':count}
    def worker(task,context):
        folder=Path(context['folder']);parts=task['id'].split('_');n=int(parts[1]);action=int(parts[2]);unit=UNITS[n]
        base='unit_%02d'%n
        if action==0:
            valid(unit,unit[2]);name=base+'-spec.json';protocol.save(folder/name,profile_data[n]);return result([name])
        if action==1:
            book=Workbook();sheet=book.active;sheet.title=unit[0][:31]
            sheet.append([unit[0]]);sheet.append(['Parameter','Value','Unit'])
            for label,value,dimension in zip(unit[3],unit[2],unit[4]):sheet.append([label,value,dimension])
            formula='='+unit[1]
            for key,row in zip('abc',range(3,6)):formula=formula.replace(key,'B'+str(row))
            sheet['A8']='Calculated result';sheet['B8']=formula;sheet['C8']=unit[5]
            sheet['A10']='Illustrative estimates. Supported inputs: 0 to 1,000,000; divisors: at least 0.000000001.'
            name=base+'-draft.xlsx';book.save(folder/name)
            assert load_workbook(folder/name).active['B8'].data_type=='f';return result([name])
        if action==2:
            count=checks(unit);name=base+'-reference.json'
            protocol.save(folder/name,{'successful_checks':count,'sample_reference':reference(unit,unit[2]),
                                      'reference_method':'rational arithmetic; independent of Excel engine'})
            return result([name],count)
        book=load_workbook(folder/(base+'-draft.xlsx'));sheet=book.active
        sheet.freeze_panes='B3'
        for row in (1,2,8):
            for cell in sheet[row][:3]:cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(color='FFFFFF',bold=True)
        for row in range(3,3+len(unit[2])):
            sheet.cell(row,2).protection=Protection(locked=False)
            denominators={node.right.id for node in ast.walk(ast.parse(unit[1],mode='eval')) if isinstance(node,ast.BinOp) and isinstance(node.op,ast.Div) and isinstance(node.right,ast.Name)}
            is_divisor='abc'[row-3] in denominators
            validation=DataValidation(type='decimal',operator='between',formula1=1e-9 if is_divisor else 0,formula2=1000000,allow_blank=False)
            validation.error='Enter a supported finite quantity between the documented bounds; divisors must be positive.'
            validation.showErrorMessage=True;sheet.add_data_validation(validation);validation.add(sheet.cell(row,2))
        sheet.protection.sheet=True
        for column,width in [('A',34),('B',22),('C',26)]:sheet.column_dimensions[column].width=width
        name=base+'-reviewed.xlsx';book.save(folder/name)
        reopened=load_workbook(folder/name).active
        assert reopened.freeze_panes=='B3' and not reopened['B3'].protection.locked
        assert reopened['B8'].protection.locked and reopened['B8'].data_type=='f'
        return result([name],4)
    def assembler(plan,context):
        folder=Path(context['folder']);book=Workbook();guide=book.active;guide.title='Guide'
        guide.append(['Capacity planning workbook','Technical internal edition'])
        guide.append(['Purpose','Transparent arithmetic estimates, not guaranteed production outcomes.'])
        guide.append(['Instructions','Edit unlocked input cells; formula cells are protected without a secret password.'])
        guide.append(['Limits','Constant rates and independent resources are assumed; validate assumptions before operational decisions.'])
        guide.append(['Accessibility scope','Named sheets, frozen input headers, wide labels and high-contrast headers; no universal accessibility certification.'])
        guide.append(['Review scope','Automated mathematical and file compatibility checks; not scientific or commercial approval.'])
        for n,unit in enumerate(UNITS):
            source_sheet=load_workbook(folder/('unit_%02d-reviewed.xlsx'%n)).active
            sheet=book.create_sheet(unit[0][:31]);sheet.freeze_panes=source_sheet.freeze_panes
            for row in source_sheet:
                for cell in row:
                    out=sheet.cell(cell.row,cell.column,cell.value)
                    if cell.has_style:
                        out.font=copy.copy(cell.font);out.fill=copy.copy(cell.fill)
                        out.border=copy.copy(cell.border);out.alignment=copy.copy(cell.alignment)
                        out.number_format=cell.number_format
                    out.protection=copy.copy(cell.protection)
            for key,value in source_sheet.column_dimensions.items():sheet.column_dimensions[key].width=value.width
            sheet.data_validations=copy.deepcopy(source_sheet.data_validations)
            sheet.protection=copy.copy(source_sheet.protection)
            sheet['A12']='Return to Guide';sheet['A12'].hyperlink="#'Guide'!A1"
            guide.append([unit[0],unit[5]]);guide.cell(guide.max_row,1).hyperlink="#'"+sheet.title+"'!A1"
        book.save(folder/'workbook.xlsx')
        office=tools['tools']['libreoffice'];env=dict(os.environ,OMP_NUM_THREADS='1',OPENBLAS_NUM_THREADS='1')
        def convert(path,kind,destination):
            response=subprocess.run([office,'-env:UserInstallation='+(folder/'lo-profile').as_uri(),
                '--headless','--convert-to',kind,'--outdir',str(destination),str(path)],capture_output=True,text=True,timeout=45,env=env)
            if response.returncode:raise RuntimeError('WORKBOOK_OFFICE_FAILED:'+response.stderr[-500:])
        convert(folder/'workbook.xlsx','ods',folder)
        recalculated=folder/'recalculated';recalculated.mkdir()
        convert(folder/'workbook.ods','xlsx',recalculated)
        os.replace(recalculated/'workbook.xlsx',folder/'workbook.xlsx')
        text=['# Capacity planning workbook','Illustrative estimates. Replace inputs with actual measurements.',
              'Edit unlocked numeric inputs. Supported range: 0 to 1,000,000; divisor minimum: 0.000000001. Values are planning estimates, not exact physical counts.',
              'Spreadsheet floating-point arithmetic is used. Sample recalculations are checked with 1e-8 relative tolerance.',
              'Binary storage units are explicit: one GiB equals 1024 MiB. These must not be mixed with decimal GB and MB.',
              'This tool is not an agronomic recommendation, financial forecast or commercial quality certificate.']
        for u in UNITS:text.extend(['## '+u[0],'Inputs: '+', '.join(u[3]),'Output unit: '+u[5],
                                   'Definition: '+u[1]+'. Sample expected result: '+str(reference(u,u[2]))])
        (folder/'user-guide.md').write_text('\n\n'.join(text),encoding='utf8')
        names=['workbook.xlsx','workbook.ods','user-guide.md']
        with zipfile.ZipFile(folder/(PRODUCT+'-internal-qa.zip'),'w',zipfile.ZIP_DEFLATED) as archive:
            for name in names:archive.write(folder/name,PRODUCT+'/'+name)
        return result(names+[PRODUCT+'-internal-qa.zip'],3)
    def reviewer(plan,context):
        folder=Path(context['folder']);book=load_workbook(folder/'workbook.xlsx',data_only=True)
        suite=unittest.TestSuite()
        def equal(unit):
            self_test=unittest.TestCase();self_test.assertTrue(math.isclose(book[unit[0][:31]]['B8'].value,reference(unit,unit[2]),rel_tol=1e-8,abs_tol=1e-8))
        for unit in UNITS:suite.addTest(unittest.FunctionTestCase(lambda unit=unit:equal(unit)))
        measured=unittest.TextTestRunner(stream=io.StringIO()).run(suite)
        if not measured.wasSuccessful() or measured.testsRun!=12:raise ValueError('WORKBOOK_INDEPENDENT_FORMULA_QA_FAILED')
        for unit in UNITS:checks(unit)
        for name in ('workbook.xlsx','workbook.ods'):
            with zipfile.ZipFile(folder/name) as archive:
                assert archive.testzip() is None
                assert not any('vbaProject' in p for p in archive.namelist())
        native=load_workbook(folder/'workbook.xlsx',data_only=False)
        assert len(native.sheetnames)==13 and (folder/'user-guide.md').stat().st_size>1500
        assert all(native[u[0][:31]]['B8'].data_type=='f' for u in UNITS)
        # Verify the remaining bounded technical criteria before recording gates.
        for unit in UNITS:
            sheet=native[unit[0][:31]]
            assert sheet.freeze_panes=='B3'
            assert sheet['B8'].protection.locked and not sheet['B3'].protection.locked
            assert sheet['A12'].hyperlink is not None
            assert sheet['A2'].fill.fgColor.rgb[-6:]=='17365D'
            assert sheet['A2'].font.color.rgb[-6:]=='FFFFFF'
            assert all(any(('B'+str(row)) in rule.sqref for rule in sheet.data_validations.dataValidation) for row in range(3,3+len(unit[2])))
        with zipfile.ZipFile(folder/(PRODUCT+'-internal-qa.zip')) as archive:
            assert archive.testzip() is None
            assert sorted(archive.namelist())==sorted(PRODUCT+'/'+name for name in ('workbook.xlsx','workbook.ods','user-guide.md'))
        # Explicit bounded automated review scope; owner acceptance remains required.
        qa={'tests_run':measured.testsRun,'successful_tests':measured.testsRun,'failures':len(measured.failures),
            'errors':len(measured.errors),'scope':'12 mathematical tools; native and LibreOffice recalculation',
            'scientific_or_commercial_approval':False}
        protocol.save(folder/'independent-review.json',qa)
        return dict(result(['independent-review.json'],12),reviewer_id='libreoffice_rational_reference_validator',
                    critical_defects=[],gates=dict.fromkeys(protocol.GATES,True),automated_review_scope=qa['scope'])
    if shutil.disk_usage(root).free<2*1024**3:raise RuntimeError('WORKBOOK_DISK_RESERVE_REQUIRED')
    state=protocol.run_attempt(root,plan,worker,assembler,reviewer)
    if state['status']!='awaiting_owner_acceptance':return state
    folder=Path(state['attempt_folder']);target.mkdir(parents=True,exist_ok=True,mode=0o700)
    archive=folder/(PRODUCT+'-internal-qa.zip');temporary=target/'.incoming.zip'
    shutil.copyfile(archive,temporary);os.replace(temporary,target/archive.name)
    qa=protocol.read(folder/'independent-review.json')
    protocol.save(target/'manufacturing_manifest.json',{'schema_version':2,'recipe':PRODUCT,'name':'Capacity planning workbook',
        'source_sha256':protocol.digest({name:protocol.file_sha(folder/name) for name in ('workbook.xlsx','workbook.ods','user-guide.md')}),'build50_source_snapshot_sha256':source,'artifact_sha256':protocol.file_sha(archive),
        'files':['workbook.xlsx','workbook.ods','user-guide.md'],'qa_status':'passed','qa_evidence':'unit_tests_passed',
        'qa_measurements':qa,'commercial_status':'private_internal_qa_pass_not_published',
        'public_listing_enabled':False,'customer_order_required':False,'owner_publication_approval':False})
    return {'status':'technical_50_stage_build_completed','completed_stages':50,'commercial_release':False}

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--factory-private',required=True);parser.add_argument('--toolchain',required=True);parser.add_argument('--worker',action='store_true')
    args=parser.parse_args();tools=protocol.read(args.toolchain)
    required={'images','documents','office_render_and_recalculation'}
    if not required.issubset(tools.get('checks',{})) or any(tools['checks'][k].get('passed') is not True for k in required):raise RuntimeError('WORKBOOK_TOOLCHAIN_NOT_VERIFIED')
    if not args.worker:
        result=subprocess.run(['nice','-n','15',tools['python'],str(Path(__file__).resolve()),'--worker',
                '--factory-private',args.factory_private,'--toolchain',args.toolchain],capture_output=True,text=True,timeout=180,
                env=dict(os.environ,OMP_NUM_THREADS='1',OPENBLAS_NUM_THREADS='1'))
        print(result.stdout,end='');print(result.stderr,end='',file=sys.stderr);sys.exit(result.returncode)
    print('FACTORY_WORKBOOK_BUILD50='+json.dumps(build(args.factory_private,tools),sort_keys=True))
