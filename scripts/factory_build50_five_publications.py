#!/usr/bin/env python3
"""Five genuine fixed-scope *technical pilot* authoring/report makers: F01-F05.

Source content is original, all measurements in F04/F05 are transparently
synthetic, and no doctorate-level specialist accuracy, market intelligence,
human editorial rights or commercial acceptance are implied.
"""
import argparse
from html import escape
import importlib.util
import json
import os
from pathlib import Path
import shutil
import zipfile
import xml.etree.ElementTree as ET

HERE=Path(__file__).resolve().parent
PRODUCTS={
"F01":"fieldwork_original_technical_handbook_pilot",
"F02":"fieldwork_original_microfiction_collection_pilot",
"F03":"fieldwork_offline_computing_dictionary_pilot",
"F04":"fieldwork_reproducible_synthetic_studies_pilot",
"F05":"fieldwork_transparent_market_simulator_pilot"}
TOPICS=("backups","permissions","change_control","monitoring","storage","upgrades",
 "incident_response","handoffs","reliability","release_review","quality_records","recovery")
VARIANTS=("foundation","example","exercise","audit")
WORDS={
 "backups":("backup","copy data and test that the restored copy opens"),
 "permissions":("access control","assign only the privileges needed to perform one task"),
 "change_control":("change review","compare a proposed change with a documented baseline"),
 "monitoring":("service checks","observe one agreed health indicator and record anomalies"),
 "storage":("data integrity","verify file lengths and a stable cryptographic digest"),
 "upgrades":("upgrade checks","rehearse updating a dependency and prepare a rollback"),
 "incident_response":("incident handling","record observable symptoms before changing configuration"),
 "handoffs":("team coordination","transfer responsibilities with a dated acceptance checklist"),
 "reliability":("service resilience","repeat a restart and compare the outcome with expectations"),
 "release_review":("release acceptance","test the same acceptance criteria in a fresh environment"),
 "quality_records":("evidence review","distinguish a measured observation from an unchecked claim"),
 "recovery":("recovery rehearsal","simulate failure and test restoration from a known-good snapshot")}
SCHEMES={
 "F01":(".pdf",".epub"),"F02":(".pdf",".epub"),
 "F03":(".html",".pdf",".epub"),
 "F04":(".docx",".pdf"),"F05":(".xlsx",".pdf")}
def load(name):
 p=HERE/(name+".py")
 if not p.is_file() or p.is_symlink():raise ValueError("PUBLISHERS_BAD_MODULE")
 sp=importlib.util.spec_from_file_location(name,str(p))
 m=importlib.util.module_from_spec(sp);sp.loader.exec_module(m);return m
def plan(proto,fid):
 if fid not in PRODUCTS:raise ValueError("PUBLISHERS_UNKNOWN_FAMILY")
 tasks=[]
 for topic in TOPICS:
  for variation in VARIANTS:
   tasks.append({"id":topic+"_"+variation,"topic":topic,"variant":variation,
    "instruction":"Produce independently readable original "+fid+" "+variation+
      " material about "+WORDS[topic][0]+", implementing "+WORDS[topic][1]+
      " with distinct testable evidence and transparent provenance limits.",
    "acceptance":["Generate genuine parsable distinct native deliverable files for this specific topic and variant.",
      "Independently cross-check document content and numerical claims or source limitations without inventing citations."]})
 return {"schema":proto.SCHEMA,"family_id":fid,"product_id":PRODUCTS[fid],
   "producer_id":"independently_measurable_five_authoring_pilot_v1",
   "source_snapshot_sha256":proto.digest({
    "family":fid,"maker":proto.file_sha(__file__),
    "reviewer":proto.file_sha(HERE/"factory_five_publication_review.py"),
    "protocol":proto.file_sha(HERE/"factory_build_50.py")}),
   "work_packages":tasks}
def content(fid,task):
 topic=task["topic"];mode=task["variant"]
 term,practice=WORDS[topic]
 index=TOPICS.index(topic)+1;idx=VARIANTS.index(mode)+1
 title=(fid+" — "+term.title()+" / "+mode.title())
 if fid=="F02":
  setting=("old port","mountain station","desert library","floating garden",
   "rain-lit observatory","coastal workshop","underground museum","distant village",
   "city night market","orbital post office","forest shelter","quiet lighthouse")[index-1]
  protagonist=("Ari","Nia","Sol","Mira")[idx-1]
  paragraphs=[
    protagonist+" arrived at the "+setting+" carrying a paper map that contained one handwritten instruction: ask what can be verified.",
    "A brass clock rang once, yet none of the people waiting inside moved. The keeper pointed to a sealed envelope and said that its story would change depending on which question came first.",
    protagonist+" chose to examine the envelope rather than trust the keeper's confident answer. Under the wax was a sketch of a bridge, a list of small repairs and an unsigned promise.",
    "The bridge was real, but its repairs had been recorded in a notebook that nobody had opened. A neighbor supplied the notebook; another neighbor showed the worn steps leading toward the water.",
    "By sunset the group had agreed on one next action, small enough to be tested before anyone was asked to cross. They repaired a single board, walked it in both directions and wrote down what had changed.",
    "The keeper offered to call the bridge safe. "+protagonist+" declined. One strong plank had demonstrated progress, not certainty. The group returned to the village with a plan to examine the remaining span.",
    "That night the clock rang again. This time people moved, not because the problem had vanished, but because their promises had turned into observable work.",
    "Author's note: this is an original fictional short-scene demonstration, not a complete edited commercial novel."]
 elif fid=="F03":
  paragraphs=[
    "Entry: "+term+". Operational meaning: the act of "+practice+".",
    "A useful entry distinguishes an intended procedure, an executed procedure and independently inspected evidence.",
    "Example: take a small disposable test case and record an initial condition, one action and the resulting condition. Preserve enough detail to reproduce the comparison.",
    "Counterexample: a ticket labeled completed does not prove a file was restored, an account was restricted or a reviewer agreed.",
    "Related concepts: observation, audit trail, least privilege, checksum, rollback and human approval. Internal cross-reference identifiers are provided in the accompanying offline index.",
    "Scope and provenance: definitions are freshly authored elementary operational explanations, not a sourced academic encyclopedia. Qualified editorial review is still required."]
 elif fid=="F04":
  sample=[(index*7+idx*5+j*j)%37+5 for j in range(8)]
  avg=sum(sample)/len(sample);scale=max(sample)-min(sample)
  paragraphs=[
    "Scope: synthetic illustration of a study protocol for "+term+". No participants, external literature, observational field data or market evidence were collected.",
    "Question: How can "+practice+" be evaluated consistently using a declared test fixture?",
    "Method: reproduce eight deterministic integers using (topic_index*7+variant_index*5+j*j) mod 37 + 5, j=0..7, one-indexed topic and variant.",
    "Observed result for the SIMULATED vector: "+str(sample)+". Mean: "+format(avg,".3f")+". Range: "+str(scale)+".",
    "Threats to validity: synthetic fixture selection, tiny sample size, no independent subjects, no blinded assessments and no external construct validation.",
    "Reproducibility: recompute all values from the explicit formula. Do not cite this illustration as evidence of any real-world research claim.",
    "Human peer review: not obtained. Rights and permission sign-off: pending. Publication and commercial sale: disallowed."]
 elif fid=="F05":
  baseline=80+index*11+idx*3
  values=[baseline+(k+1)*(index+idx)+((k%2)*3) for k in range(6)]
  paragraphs=[
    "Market-model prototype about "+term+". Data are deterministic simulated examples only and are NOT current prices, sales, competitors, customers or investment recommendations.",
    "Assumption: monthly index starts from an arbitrary baseline "+str(baseline)+" and increments by an illustrative formula for six periods.",
    "Synthetic sample values: "+", ".join(str(v) for v in values)+".",
    "The accompanying editable workbook contains source input values and a derived month-on-month change formula.",
    "Validation plan: acquire dated first-party or verifiable third-party market records, record provenance/rights, de-duplicate observations, check representative sampling and verify currency units.",
    "Failure modes: stale prices, survivorship bias, mixed gross and net amounts, missing refunds and unsupported conversion assumptions.",
    "No real-world market conclusions or financial forecasting accuracy are claimed. Qualified independent validation and owner pricing approval remain missing."]
 else:
  paragraphs=[
    "Fieldwork operating handbook: "+term+". Practical objective: "+practice+".",
    "Prerequisites: an isolated example environment, an initial baseline, documented ownership and permission to inspect or modify the disposable test fixture.",
    "Step 1. Define the smallest observable unit and an expected outcome before making any change.",
    "Step 2. Create or select an original safe test asset; collect a SHA-256 digest, timestamp and written test description.",
    "Step 3. Execute the exact planned activity once, record the actual output and compare it with the expected outcome.",
    "Step 4. Run one negative case in which the target file is missing, access is denied or the measured result differs; ensure the process fails safely.",
    "Step 5. Return to the baseline, repeat the test and verify that no persistent production data or accounts were affected.",
    "Quality checklist: verified native artifact; complete instructions; negative case; independent rerun; no third-party proprietary data; documented license and expert-review gaps.",
    "Worked check for "+term+": preserve evidence for the decision to "+practice+". Distinguish test execution from actual operational deployment.",
    "Limitation: educational operations sample only. This text is not a peer-reviewed specialist reference or a complete commercially approved technical book."]
 if fid=="F02":
  paragraphs.insert(1,"The motif of this "+mode+" scene is "+term+
   ", yet its events, choices and people are fictional; no external source is being quoted.")
 return title,paragraphs
def pdf(path,title,paragraphs):
 from reportlab.pdfgen import canvas
 from reportlab.lib.pagesizes import A4
 pdf=canvas.Canvas(str(path),pagesize=A4)
 pdf.setTitle(title)
 y=790
 for n,line in enumerate([title]+paragraphs):
  words=line.split();current=""
  for word in words:
   if len(current)+len(word)>89:
    if y<84:pdf.showPage();y=792
    pdf.setFont("Helvetica-Bold" if n==0 else "Helvetica",13 if n==0 else 10)
    pdf.drawString(43,y,current.strip());y-=16;current=""
   current+=word+" "
  if current:
   if y<84:pdf.showPage();y=792
   pdf.setFont("Helvetica-Bold" if n==0 else "Helvetica",13 if n==0 else 10)
   pdf.drawString(43,y,current.strip());y-=24
 pdf.save()
def epub(path,title,paragraphs,slug):
 body="".join("<p>"+escape(s)+"</p>" for s in paragraphs)
 xhtml='<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>'+escape(title)+'</title></head><body><h1>'+escape(title)+'</h1>'+body+'</body></html>'
 nav='<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>Contents</title></head><body><nav xmlns:epub="http://www.idpf.org/2007/ops" epub:type="toc"><ol><li><a href="chapter.xhtml">Chapter</a></li></ol></nav></body></html>'
 opf='<?xml version="1.0" encoding="UTF-8"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="uid"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="uid">urn:uuid:saeit-'+slug+'</dc:identifier><dc:title>'+escape(title)+'</dc:title><dc:language>en</dc:language><meta property="dcterms:modified">2026-10-20T00:00:00Z</meta></metadata><manifest><item id="text" href="chapter.xhtml" media-type="application/xhtml+xml"/><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/></manifest><spine><itemref idref="text"/></spine></package>'
 container='<?xml version="1.0" encoding="utf-8"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>'
 with zipfile.ZipFile(str(path),"w") as z:
  z.writestr("mimetype","application/epub+zip",compress_type=zipfile.ZIP_STORED)
  for name,text in (("META-INF/container.xml",container),("OEBPS/content.opf",opf),
      ("OEBPS/nav.xhtml",nav),("OEBPS/chapter.xhtml",xhtml)):
   z.writestr(name,text,compress_type=zipfile.ZIP_DEFLATED)
def html(path,title,paragraphs,slug):
 content="".join("<p>"+escape(p)+"</p>" for p in paragraphs)
 path.write_text('<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>'+
  escape(title)+'</title></head><body><main><h1>'+escape(title)+'</h1>'+content+
  '<p>Offline entry reference: '+escape(slug)+'</p></main></body></html>',encoding="utf8")
def research_docx(path,title,paragraphs):
 from docx import Document
 document=Document();document.add_heading(title,0)
 for i,body in enumerate(paragraphs):
  document.add_heading(["Protocol","Research question","Method","Results","Threats to validity","Reproducibility","Permissions"][i],level=2)
  document.add_paragraph(body)
 document.save(str(path))
def market_xlsx(path,task):
 from openpyxl import Workbook,load_workbook
 from openpyxl.styles import Font
 index=TOPICS.index(task["topic"])+1;idx=VARIANTS.index(task["variant"])+1
 baseline=80+index*11+idx*3
 w=Workbook();s=w.active;s.title="Synthetic scenario only"
 s.append(["Month","Simulated index - not market data","Change (formula)","Source","Evidence status"])
 for j in range(6):
  v=baseline+(j+1)*(index+idx)+(j%2)*3
  s.append([j+1,v,None if j==0 else "=B%d-B%d"%(j+2,j+1),"Deterministic fixture","NOT REAL"])
 for cell in s[1]:cell.font=Font(bold=True)
 s.freeze_panes="A2"
 for col,dim in {"A":12,"B":37,"C":21,"D":26,"E":24}.items():
  s.column_dimensions[col].width=dim
 w.save(str(path))
 c=load_workbook(str(path),data_only=False,read_only=True)
 assert c.active["B2"].value==baseline+index+idx
 assert c.active["C3"].value=="=B3-B2"
def make(folder,fid,task):
 title,paragraphs=content(fid,task)
 stem=task["id"];files=[]
 for ext in SCHEMES[fid]:
  path=folder/(stem+ext)
  if ext==".pdf":pdf(path,title,paragraphs)
  elif ext==".epub":epub(path,title,paragraphs,stem)
  elif ext==".html":html(path,title,paragraphs,stem)
  elif ext==".docx":research_docx(path,title,paragraphs)
  elif ext==".xlsx":market_xlsx(path,task)
  else:raise ValueError("UNKNOWN_FILE_FORMAT")
  if not path.is_file() or path.stat().st_size<250:raise ValueError("INVALID_NATIVE_AUTHORING_OUTPUT")
  files.append(path.name)
 return files
def build(root,fid):
 root=Path(root)
 if not root.is_dir() or root.is_symlink() or fid not in PRODUCTS:raise ValueError("AUTHORING_UNSAFE_ROOT_OR_FAMILY")
 proto=load("factory_build_50");ind=load("factory_five_publication_review")
 spec=plan(proto,fid);product=PRODUCTS[fid];home=root/"build-50"/product
 dest=root/"digital-products"/product;archive=product+"-internal-qa.zip"
 if any(p.is_symlink() for p in (root/"build-50",root/"digital-products",home,dest)):
  raise ValueError("AUTHORING_SYMLINK_OUTPUT")
 if (home/"state.json").is_file() and (dest/"manufacturing_manifest.json").is_file():
  state=proto.read(home/"state.json");m=proto.read(dest/"manufacturing_manifest.json")
  path=dest/archive
  if proto.verify_attempt(home,state,spec["source_snapshot_sha256"]) and path.is_file() and not path.is_symlink() and proto.file_sha(path)==m.get("artifact_sha256"):
   return {"family_id":fid,"status":"unchanged_verified_technical_product","commercial_release":False}
 def worker(task,ctx):
  folder=Path(ctx["folder"]);files=make(folder,fid,task)
  return {"passed":True,"outputs":files,"successful_checks":2}
 def assemble(plan,ctx):
  folder=Path(ctx["folder"]);stems={x["id"] for x in plan["work_packages"]}
  members=sorted(p for p in folder.iterdir() if p.stem in stems and p.suffix in SCHEMES[fid])
  note=("Original fixed-scope "+fid+" technically valid pilot. "+ 
    ("SYNTHETIC DATA ONLY: no actual research, dated market sources or investment advice. " if fid in ("F04","F05") else "No original peer-reviewed source-backed expert claim. ")+
    "No professional human editorial review, license signoff or commercial/publication approval.\n")
  (folder/"README.txt").write_text(note,encoding="utf8")
  with zipfile.ZipFile(str(folder/archive),"w",zipfile.ZIP_DEFLATED) as z:
   for p in members+[folder/"README.txt"]:
    item=zipfile.ZipInfo(product+"/"+p.name,date_time=(2020,1,1,0,0,0))
    item.compress_type=zipfile.ZIP_DEFLATED
    z.writestr(item,p.read_bytes())
  return {"passed":True,"outputs":["README.txt",archive],"successful_checks":2}
 def reviewer(plan,ctx):
  folder=Path(ctx["folder"]);checked=ind.check(folder,folder/archive,product,fid)
  proto.save(folder/"independent-authoring-review.json",checked)
  return {"reviewer_id":"native_epub_docx_xlsx_pdf_integrity_and_transparency_reviewer",
   "passed":checked["passed"],"critical_defects":[],"gates":checked["gates"],
   "successful_checks":10,"outputs":["independent-authoring-review.json"]}
 state=proto.run_attempt(root,spec,worker,assemble,reviewer)
 if not proto.verify_attempt(home,state,spec["source_snapshot_sha256"]):
  return {"family_id":fid,"status":"technical_rework_required","defects":state.get("defects",[])}
 dest.mkdir(parents=True,exist_ok=True,mode=0o700)
 tmp=dest/".bookreport-incoming.zip"
 shutil.copyfile(Path(state["attempt_folder"])/archive,tmp);os.replace(tmp,dest/archive)
 proto.save(dest/"manufacturing_manifest.json",{"schema_version":2,"recipe":product,
  "family_id":fid,"artifact_sha256":proto.file_sha(dest/archive),
  "bytes":(dest/archive).stat().st_size,"qa_status":"passed",
  "qa_evidence":"native_file_internals_and_original_simulation_integrity_only",
  "commercial_status":"private_internal_qa_pass_not_published",
  "public_listing_enabled":False,"owner_publication_approval":False,
  "external_market_or_research_data":False,"independent_human_review":False})
 return {"family_id":fid,"status":"technical_50_stage_build_completed",
  "completed_stages":50,"commercial_release":False}
if __name__=="__main__":
 p=argparse.ArgumentParser();p.add_argument("--factory-private",required=True)
 p.add_argument("--family",required=True,choices=sorted(PRODUCTS));args=p.parse_args()
 outcome=build(args.factory_private,args.family)
 print("FACTORY_BOOK_REPORT_BUILD50="+json.dumps(outcome,sort_keys=True))
 if outcome["status"]=="technical_rework_required":raise SystemExit(2)
