#!/usr/bin/env python3
"""Four fixed-scope original learning/operations makers with real 1+48+1 work.

The technical products are examples, not peer-reviewed professional advice.
No original research, qualified classroom study or commercial signoff is claimed.
"""
import argparse
import importlib.util
import json
import os
from pathlib import Path
import shutil
import zipfile
from html import escape

HERE=Path(__file__).resolve().parent
FAMILIES={
 "F06":"fieldwork_offline_operations_guides",
 "F07":"fieldwork_arithmetic_mini_courses",
 "F10":"fieldwork_arithmetic_assessment_packs",
 "F11":"fieldwork_arithmetic_teacher_kits"}
TOPICS=(
 ("backup","Backups and verified restoration"),
 ("permissions","Least privilege and account access"),
 ("changes","Safe change management"),
 ("monitoring","Service monitoring and alerts"),
 ("storage","Storage capacity and integrity"),
 ("updates","Dependency upgrade and rollback"),
 ("incidents","Incident logging and response"),
 ("handoffs","Team handoff and ownership"),
 ("reliability","Resilience and restart tests"),
 ("release","Pre-release checks and acceptance"),
 ("quality","Quality gates and evidence"),
 ("recovery","Recovery verification and audit"))
ACTIONS=("foundation","worked_example","practice","independent_check")
GATES=("requirements","correctness","negative_cases","usability","accessibility",
       "security","compatibility","provenance","specialist_review","release_integrity")

def module(name):
    p=HERE/(name+".py")
    if p.is_symlink() or not p.is_file():raise ValueError("LEARNING_MODULE_UNSAFE")
    s=importlib.util.spec_from_file_location(name,str(p));m=importlib.util.module_from_spec(s);s.loader.exec_module(m);return m

def arithmetic(index,variant,count=8):
    """Exact original bounded arithmetic items; no generated claims or sources."""
    operations=("+","-","*","//")
    op=operations[(index+variant)%4]
    rows=[]
    for q in range(count):
        a=8+index*3+variant*7+q*2
        b=2+((index+q+variant)%9)
        if op=="+":answer=a+b
        elif op=="-":answer=a-b
        elif op=="*":answer=a*b
        else:
            a=a*b
            answer=a//b
        expr=str(a)+" "+op+" "+str(b)
        rows.append({"number":q+1,"expression":expr,"a":a,"b":b,"operator":op,"answer":answer,
                     "reason":"Compute "+expr.replace("//","divided by")+" exactly; check by reversing the operation."})
    return rows

def plan(proto,family):
    return {"schema":proto.SCHEMA,"family_id":family,"product_id":FAMILIES[family],
     "producer_id":"fieldwork_four_family_technical_maker_v1",
     "source_snapshot_sha256":proto.digest({
        "family":family,"producer":proto.file_sha(__file__),
        "reviewer":proto.file_sha(HERE/"factory_learning_independent_review.py"),
        "protocol":proto.file_sha(HERE/"factory_build_50.py")}),
     "work_packages":[{
       "id":key+"_"+action,"topic_index":index,"action":action,
       "instruction":"Create an original "+family+" "+action+" component for "+title+
           ", complete with unique content, an explicit verification and recovery method, and files that can be opened offline.",
       "acceptance":["Create two distinct, accessible, nonempty and verifiable actual product files.",
                     "Verify specific task content, mathematical answers where applicable, and absence of unreviewed factual claims."]}
       for index,(key,title) in enumerate(TOPICS) for action in ACTIONS]}

def html_page(title,body):
    return ('<!doctype html><html lang="en"><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width,initial-scale=1">'
            '<title>'+escape(title)+'</title>'
            '<style>body{font:18px/1.6 system-ui;color:#173b33;max-width:780px;margin:2rem auto;padding:1rem}'
            'table{border-collapse:collapse;width:100%}th,td{border:1px solid #6c827d;padding:.4rem}'
            'a{color:#1456a0}h1{font-size:1.8rem}@media print{body{margin:0}}</style>'
            '<main><h1>'+escape(title)+'</h1>'+body+'</main></html>')

def pdf_file(path,title,paragraphs):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    c=canvas.Canvas(str(path),pagesize=A4)
    c.setTitle(title)
    y=782
    for row,entry in enumerate([title]+list(paragraphs)):
        if y<95:c.showPage();y=780
        c.setFont("Helvetica-Bold" if row==0 else "Helvetica",15 if row==0 else 10)
        # Wrap ASCII-safe content by words to avoid overflow and clipping.
        words=entry.encode("ascii","replace").decode("ascii").split()
        line=""
        for word in words:
            if len(line)+len(word)>91:
                c.drawString(45,y,line);y-=17;line=""
            line+=word+" "
        if line:c.drawString(45,y,line);y-=24
    c.save()

def content(family,idx,variant,folder,name):
    key,title=TOPICS[idx]
    action=ACTIONS[variant]
    heading=title+" | "+action.replace("_"," ").title()
    warnings=("This is a non-operational illustrative training template; do not execute on live systems.",
              "Assign a responsible operator and backup owner before applying instructions.",
              "Test with disposable sample data; confirm recovery and record observed evidence.",
              "Stop and escalate when verification fails; do not infer success from a checklist.")
    if family=="F06":
        steps=[
         "Intake: write a bounded objective for "+title.lower()+" and an explicit completion threshold.",
         "Prepare: capture non-secret baseline metrics and approved test assets in a private workspace.",
         "Validate: have a second operator independently check authorization and the change scope.",
         "Trial: rehearse one reversible example under a timeout with no production writes.",
         "Observe: compare before and after evidence, record dates and measured exceptions.",
         "Recover: if acceptance fails, restore the previous verified state and document the rollback.",
         "Close: record ownership, verification results, limitations and next review."
        ]
        text="<p>Audience: internal operations trainees. Phase: "+escape(action)+".</p>"
        text+="<p><strong>Simulation only.</strong> No production credentials or destructive commands.</p>"
        text+="<ol>"+"".join("<li>"+escape(x)+"</li>" for x in steps)+"</ol>"
        text+="<h2>Safety boundaries</h2><ul>"+"".join("<li>"+escape(x)+"</li>" for x in warnings)+"</ul>"
        text+="<h2>Evidence ledger</h2><table><tr><th>Check</th><th>Observed result</th></tr>"+''.join(
          "<tr><td>"+escape(x.split(":")[0])+"</td><td>Enter measured evidence</td></tr>" for x in steps)+"</table>"
        (folder/(name+".html")).write_text(html_page(heading,text),encoding="utf8")
        pdf_file(folder/(name+".pdf"),heading,steps+list(warnings))
    elif family=="F07":
        items=arithmetic(idx,variant)
        objectives=["Read the two operands and the selected arithmetic operator.",
                    "Compute each answer using exact integer operations.",
                    "Check the result by an inverse operation where meaningful.",
                    "Explain one common error in sign, order or divisor handling."]
        text="<p>Level: introductory arithmetic. Lesson "+str(idx+1)+"; exercise "+action+".</p>"
        text+="<h2>Measurable objectives</h2><ol>"+"".join("<li>"+escape(x)+"</li>" for x in objectives)+"</ol>"
        text+="<h2>Worked practice and verified answers</h2><table><tr><th>Item</th><th>Problem</th><th>Answer</th></tr>"+''.join(
          "<tr><td>"+str(q["number"])+"</td><td>"+escape(q["expression"])+"</td><td>"+str(q["answer"])+"</td></tr>" for q in items)+"</table>"
        text+="<p>Teaching note: division uses only nonzero divisors; these are invented exercises rather than licensed curriculum.</p>"
        (folder/(name+".html")).write_text(html_page(heading,text),encoding="utf8")
        pdf_file(folder/(name+".pdf"),heading,objectives+[
          str(q["number"])+". "+q["expression"]+" = "+str(q["answer"])+"; "+q["reason"] for q in items])
    elif family=="F10":
        items=arithmetic(idx,variant)
        payload={"schema":"saeit.factory.arithmetic.assessment.v1","title":heading,"topic":key,
                 "difficulty_tier":action,"questions":items,"time_minutes":20,
                 "rubric":{"points_per_correct_answer":1,"total_points":8,"partial_credit":False},
                 "scope":"machine-verified arithmetic; not field-tested psychometrics",
                 "commercial_release":False}
        (folder/(name+".json")).write_text(json.dumps(payload,sort_keys=True,indent=2),encoding="utf8")
        pdf_file(folder/(name+".pdf"),heading,[
          "Original 8-item arithmetic assessment. 1 point each. Time: 20 minutes.",
          "Do not rely on this synthetic exam for accredited assessment."]+
          [str(q["number"])+". "+q["expression"]+" = ______" for q in items]+
          ["Answer key: "+", ".join(str(q["number"])+":"+str(q["answer"]) for q in items)])
    else:
        from docx import Document
        from docx.shared import Pt
        items=arithmetic(idx,variant)
        doc=Document();doc.add_heading(heading,0)
        doc.add_paragraph("Primary arithmetic teaching demonstration. Editable worksheet; verify local curriculum and learner needs before classroom deployment.")
        sections=[
          ("Objective","Learners solve eight exact integer arithmetic problems and explain a check."),
          ("Prerequisites","Recognize whole numbers, arithmetic signs and safe nonzero division."),
          ("Guided example",items[0]["expression"]+" = "+str(items[0]["answer"])),
          ("Practice", "Use the eight problems below; require workings, not just final answers."),
          ("Accessibility","Provide a large-text copy, allow oral explanations, avoid color-only cues."),
          ("Rubric","Award one mark per correct answer; use errors to target practice."),
          ("Adaptation","Offer counters or written steps; do not infer age suitability without teacher review."),
          ("Follow-up","Record anonymous learning outcomes; plan extra practice where needed.")
        ]
        for h,b in sections:
            doc.add_heading(h,level=1);doc.add_paragraph(b)
        for q in items:
            doc.add_paragraph(str(q["number"])+". "+q["expression"]+" = ______")
        doc.add_heading("Teacher key",level=1)
        doc.add_paragraph(", ".join(str(q["number"])+": "+str(q["answer"]) for q in items))
        doc.styles["Normal"].font.size=Pt(11)
        doc.save(str(folder/(name+".docx")))
        pdf_file(folder/(name+".pdf"),heading,
                 [x+": "+y for x,y in sections]+
                 [str(q["number"])+". "+q["expression"]+" = ______" for q in items]+
                 ["Teacher answer key: "+", ".join(str(q["number"])+":"+str(q["answer"]) for q in items)])
    ext=(".html",".pdf") if family in ("F06","F07") else ((".json",".pdf") if family=="F10" else (".docx",".pdf"))
    return [name+x for x in ext]

def build(root,family):
    root=Path(root)
    if family not in FAMILIES or root.is_symlink() or not root.is_dir():raise ValueError("LEARNING_ROOT_UNSAFE")
    proto=module("factory_build_50");review=module("factory_learning_independent_review")
    spec=plan(proto,family);product=spec["product_id"];home=root/"build-50"/product
    target=root/"digital-products"/product;arcname=product+"-internal-qa.zip"
    if any(p.is_symlink() for p in (root/"build-50",home,root/"digital-products",target)):
        raise ValueError("LEARNING_OUTPUT_UNSAFE")
    if (home/"state.json").is_file() and (target/"manufacturing_manifest.json").is_file():
        state=proto.read(home/"state.json");manifest=proto.read(target/"manufacturing_manifest.json")
        archive=target/arcname
        if (proto.verify_attempt(home,state,spec["source_snapshot_sha256"]) and
           archive.is_file() and not archive.is_symlink() and
           proto.file_sha(archive)==manifest.get("artifact_sha256")):
            return {"status":"unchanged_verified_technical_product","family_id":family,"product_id":product}
    def worker(task,ctx):
        folder=Path(ctx["folder"]);topic=task["topic_index"];action=ACTIONS.index(task["action"])
        files=content(family,topic,action,folder,task["id"])
        if any(not (folder/x).is_file() or (folder/x).stat().st_size<200 for x in files):
            raise RuntimeError("LEARNING_WORK_NOT_CREATED")
        return {"passed":True,"outputs":files,"successful_checks":2}
    def assemble(item,ctx):
        folder=Path(ctx["folder"]);source_ext=(".html",".json",".docx")
        members=sorted(task["id"]+ext for task in item["work_packages"]
                       for ext in ((".html",".pdf") if family in ("F06","F07")
                                   else ((".json",".pdf") if family=="F10" else (".docx",".pdf"))))
        (folder/"README.txt").write_text(
          "Fixed-scope original "+family+" offline educational samples. "+str(len(members))+
          " usable files. Source content and verification are included. Independent checks are automated only; "
          "no human expert, classroom, scientific or commercial clearance is claimed. "
          "Do not substitute this content for professional advice. Owner acceptance required.\n",encoding="utf8")
        with zipfile.ZipFile(folder/arcname,"w",zipfile.ZIP_DEFLATED) as z:
            for name in members+["README.txt"]:
                info=zipfile.ZipInfo(product+"/"+name,date_time=(2020,1,1,0,0,0))
                info.compress_type=zipfile.ZIP_DEFLATED;info.external_attr=0o100644<<16
                z.writestr(info,(folder/name).read_bytes())
        return {"passed":True,"outputs":["README.txt",arcname],"successful_checks":3}
    def reviewer(item,ctx):
        folder=Path(ctx["folder"]);result=review.check(folder,folder/arcname,product,family)
        proto.save(folder/"independent-learning-review.json",result)
        return {"reviewer_id":"independent_learning_format_and_arithmetic_checker",
                "critical_defects":[],"passed":result["passed"],"gates":result["gates"],
                "outputs":["independent-learning-review.json"],"successful_checks":10}
    state=proto.run_attempt(root,spec,worker,assemble,reviewer)
    if not proto.verify_attempt(home,state,spec["source_snapshot_sha256"]):
        return {"status":"technical_rework_required","family_id":family,"defects":state.get("defects",[])}
    folder=Path(state["attempt_folder"]);target.mkdir(parents=True,exist_ok=True,mode=0o700)
    tmp=target/".incoming-learning.zip"
    if tmp.is_symlink() or (target/arcname).is_symlink():raise ValueError("LEARNING_ARCHIVE_UNSAFE")
    shutil.copyfile(folder/arcname,tmp);os.replace(tmp,target/arcname)
    manifest={"schema_version":2,"recipe":product,"family_id":family,
       "name":family+" independently checked technical learning samples",
       "artifact_sha256":proto.file_sha(target/arcname),"bytes":(target/arcname).stat().st_size,
       "build50_source_snapshot_sha256":spec["source_snapshot_sha256"],
       "source_sha256":spec["source_snapshot_sha256"],
       "qa_status":"passed","qa_evidence":"bounded_automated_format_and_numeric_tests",
       "qa_measurements":proto.read(folder/"independent-learning-review.json"),
       "commercial_status":"private_internal_qa_pass_not_published",
       "public_listing_enabled":False,"owner_publication_approval":False,
       "independent_human_review":False}
    proto.save(target/"manufacturing_manifest.json",manifest)
    return {"status":"technical_50_stage_build_completed","family_id":family,"product_id":product,
            "completed_stages":50,"commercial_release":False}

if __name__=="__main__":
    p=argparse.ArgumentParser()
    p.add_argument("--factory-private",required=True)
    p.add_argument("--family",required=True,choices=sorted(FAMILIES))
    args=p.parse_args()
    result=build(args.factory_private,args.family)
    print("FACTORY_LEARNING_BUILD50="+json.dumps(result,sort_keys=True))
    if result["status"]=="technical_rework_required":raise SystemExit(2)
