import importlib.util
from pathlib import Path
import tempfile
import unittest
import zipfile
import re
import xml.etree.ElementTree as ET
ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('workbook50',ROOT/'scripts/factory_build50_workbook.py')
producer=importlib.util.module_from_spec(spec);spec.loader.exec_module(producer)

class Workbook50Tests(unittest.TestCase):
    def test_restoring_frozen_headers_preserves_cached_formula_values(self):
        try:from openpyxl import Workbook,load_workbook
        except ImportError:self.skipTest('openpyxl required')
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'book.xlsx';book=Workbook();book.active.freeze_panes='B3';book.active['B8']='=2+2';book.save(path)
            views=producer.snapshot_views(path);temporary=Path(td)/'converted.xlsx'
            with zipfile.ZipFile(path) as source,zipfile.ZipFile(temporary,'w') as target:
                for info in source.infolist():
                    data=source.read(info.filename)
                    if info.filename=='xl/worksheets/sheet1.xml':
                        data=producer.VIEW_PATTERN.sub(b'<sheetViews><sheetView workbookViewId="0"/></sheetViews>',data)
                        xml=ET.fromstring(data)
                        cached=xml.find('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c[@r="B8"]/{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v')
                        self.assertIsNotNone(cached);cached.text='4'
                        data=ET.tostring(xml)
                    target.writestr(info,data)
            self.assertIsNone(load_workbook(temporary).active.freeze_panes)
            self.assertEqual(load_workbook(temporary,data_only=True).active['B8'].value,4)
            producer.restore_views(temporary,views)
            self.assertEqual(load_workbook(temporary).active.freeze_panes,'B3')
            self.assertEqual(load_workbook(temporary,data_only=True).active['B8'].value,4)
    def test_package_members_follow_the_warehouse_order(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);names=['workbook.xlsx','workbook.ods','user-guide.md']
            for name in names:(root/name).write_bytes(('TEST FIXTURE '+name).encode())
            archive=producer.package(root,names)
            with zipfile.ZipFile(archive) as file:
                self.assertEqual(file.namelist(),[producer.PRODUCT+'/'+n for n in sorted(names)])
                self.assertIsNone(file.testzip())
    def test_twelve_reference_models_reject_invalid_inputs(self):
        self.assertEqual(len(producer.UNITS),12)
        for unit in producer.UNITS:
            self.assertGreaterEqual(producer.checks(unit),7)
            self.assertGreaterEqual(producer.reference(unit,unit[2]),0)
    def test_storage_formulas_use_explicit_binary_units(self):
        storage=next(u for u in producer.UNITS if u[0]=='Replicated storage')
        backup=next(u for u in producer.UNITS if u[0]=='Backup duration')
        self.assertEqual(storage[5],'GiB')
        self.assertIn('MiB/file',storage[4])
        self.assertEqual(backup[4],['GiB','MiB/second'])
    def test_unsupported_large_inputs_are_rejected(self):
        for unit in producer.UNITS:
            values=list(unit[2]);values[0]=1000001
            with self.assertRaises(ValueError):producer.reference(unit,values)
    def test_zero_divisors_are_rejected(self):
        for unit in producer.UNITS:
            if '/b' in unit[1]:
                values=list(unit[2]);values[1]=0
                with self.assertRaises(ValueError):producer.reference(unit,values)
    def test_actual_48_work_packages_retained_when_assembly_tool_is_missing(self):
        try:import openpyxl
        except ImportError:self.skipTest('openpyxl required for worksheet integration')
        with tempfile.TemporaryDirectory() as td:
            result=producer.build(td,{'tools':{'libreoffice':'/nonexistent/factory-office'}})
            self.assertEqual(result['completed_stages'],49)
            self.assertEqual(result['defects'][0]['failed_stage'],50)
            self.assertEqual(result['status'],'rework_required')
            folder=Path(result['attempt_folder'])
            self.assertIn('FileNotFoundError',result['defects'][0]['reason'])
            from openpyxl import load_workbook
            assembled=load_workbook(folder/'workbook.xlsx')
            for unit in producer.UNITS:
                sheet=assembled[unit[0][:31]]
                self.assertEqual(sheet['A2'].fill.fgColor.rgb[-6:],'17365D')
                self.assertEqual(sheet['A2'].font.color.rgb[-6:],'FFFFFF')
            self.assertEqual(len(list(folder.glob('*-reviewed.xlsx'))),12)
            self.assertEqual(len(list(folder.glob('*-work.json'))),48)
            self.assertFalse((Path(td)/'digital-products/capacity_planning_workbook').exists())
