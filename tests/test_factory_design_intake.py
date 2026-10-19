import hashlib
import hmac
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[1]
def module(name):
    spec=importlib.util.spec_from_file_location(name,ROOT/'scripts'/(name+'.py'))
    x=importlib.util.module_from_spec(spec);spec.loader.exec_module(x);return x
intake=module('factory_design_intake')
designer=module('factory_recipe_designer')
maker=module('factory_product_maker')

class DesignIntakeTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
    def record(self,origin,ident='need_01',blueprint='inventory_coverage',family='F28'):
        p=self.root/'source-evidence'/ident;p.parent.mkdir(exist_ok=True)
        p.write_text('TEST FIXTURE: explicit mathematical need, not a real customer or market observation')
        record={'id':ident,'origin':origin,'design_brief':{'schema':intake.SCHEMA,
                'family_id':family,'blueprint':blueprint},
                'evidence':[{'path':str(p.relative_to(self.root)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}]}
        if origin in ('customer_request','market_demand'):
            key=b'UNIT_TEST_INDEPENDENT_REVIEW_KEY_32_BYTES'
            k=self.root/'design-source-review.key';k.write_bytes(key);k.chmod(0o600)
            record['review_signature']=hmac.new(key,intake.digest({'origin':origin,'record':record}).encode(),'sha256').hexdigest()
        return record
    def snapshot(self,old,records):
        path=self.root/'production-order-inputs'/(old+'.json')
        designer.save(path,{'schema':'saeit.factory.production-instructions.v1','items':records})
    def test_all_four_channels_traceable_and_idempotent(self):
        for i,(old,origin) in enumerate(intake.CHANNELS.items()):
            self.snapshot(old,[self.record(origin,'need_%d'%i)])
        first=intake.cycle(self.root)
        self.assertEqual(first['enqueued'],4)
        second=intake.cycle(self.root)
        self.assertEqual(second['enqueued'],0)
        self.assertEqual(sum(x['replayed'] for x in second['channels'].values()),4)
        self.assertEqual(len(list((self.root/'design-intake').glob('*.json'))),4)
        self.assertEqual(set(x['origin'] for x in first['latest_results']),intake.ORIGINS)
    def test_source_hash_tampering_never_reaches_designer(self):
        r=self.record('site_need');r['evidence'][0]['sha256']='a'*64
        self.snapshot('site-topic-snapshot',[r])
        result=intake.cycle(self.root)
        self.assertEqual(result['enqueued'],0)
        self.assertIn('HASH_MISMATCH',result['latest_results'][0]['reason'])
    def test_market_or_customer_without_independent_signature_stays_blocked(self):
        for old,origin in (('market-research-snapshot','market_demand'),('customer-intake-snapshot','customer_request')):
            record=self.record(origin,origin);record.pop('review_signature')
            self.snapshot(old,[record])
        result=intake.cycle(self.root)
        self.assertEqual(result['enqueued'],0)
        self.assertEqual(sum(x['blocked'] for x in result['channels'].values()),2)
    def test_unsupported_family_and_prose_without_engineering_are_preserved(self):
        record=self.record('site_need',family='F12')
        self.snapshot('site-topic-snapshot',[record])
        result=intake.cycle(self.root,now=0)
        self.assertIn('ADAPTER_NOT_IMPLEMENTED',result['latest_results'][0]['reason'])
        self.assertEqual(len(result['capabilities']),40)
        self.assertEqual(intake.cycle(self.root,now=1)['processed'],0)
        self.assertEqual(intake.cycle(self.root,now=61)['processed'],1)
    def test_two_high_level_needs_become_distinct_tested_real_packages(self):
        self.snapshot('site-topic-snapshot',[self.record('site_need','need_coverage','inventory_coverage'),
                                            self.record('site_need','need_water','water_volume')])
        self.assertEqual(intake.cycle(self.root)['enqueued'],2)
        results=[]
        for _ in range(2):
            d=designer.cycle(self.root)
            results.extend(d['actions'])
        self.assertEqual(len({x['result']['functional_sha256'] for x in results}),2)
        recipes=self.root/'generated-product-recipes'
        for recipe in sorted(recipes.iterdir()):
            result=maker.manufacture(recipe.name,recipes,self.root/'products')
            self.assertTrue(result['changed'])
            self.assertFalse(maker.manufacture(recipe.name,recipes,self.root/'products')['changed'])
            manifest=json.loads((self.root/'products'/recipe.name/'manufacturing_manifest.json').read_text())
            self.assertEqual(manifest['qa_measurements']['successful_tests'],3)
    def test_independent_rational_oracles_match_generated_programs_on_unseen_inputs(self):
        for key,(_,_,fields,_,oracle) in intake.BLUEPRINTS.items():
            spec=intake.expand({'schema':intake.SCHEMA,'blueprint':key},'factory_invention')
            result=designer.compile_recipe(spec,self.root/'recipes')
            path=self.root/'recipes'/result['recipe']/'calculator.py'
            s=importlib.util.spec_from_file_location('calculation_'+key,path)
            tool=importlib.util.module_from_spec(s);s.loader.exec_module(tool)
            from fractions import Fraction
            for a in (2,7,41):
                for b in (3,11,53):
                    values=[a,b,17][:len(fields)]
                    expected=oracle(*[Fraction(x) for x in values])
                    actual=Fraction(tool.calculate(dict(zip([x[0] for x in fields],values)))['result'])
                    self.assertLessEqual(abs(expected-actual),Fraction('1e-12')*max(1,abs(expected)))
    def test_real_xlsx_adapter_is_not_only_a_registry_entry(self):
        spec=intake.expand({'schema':intake.SCHEMA,'family_id':'F27','blueprint':'unit_margin'},'factory_invention')
        result=designer.compile_recipe(spec,self.root/'recipes')
        source=self.root/'recipes'/result['recipe']
        from openpyxl import load_workbook
        book=load_workbook(source/'calculator.xlsx')
        self.assertIn('Calculator',book.sheetnames)
        evidence={};maker.verify_recipe(source,evidence=evidence)
        self.assertEqual(evidence['successful_tests'],4)
        self.assertEqual(result['family_id'],'F27')
    def test_supply_is_internal_and_does_not_invent_customer_or_market_demand(self):
        supply=ROOT/'factory_config/internal_design_supply.json'
        first=intake.cycle(self.root,supply=supply)
        second=intake.cycle(self.root,supply=supply)
        self.assertEqual(first['enqueued']+second['enqueued'],15)
        self.assertEqual(first['channels']['customer_request']['accepted'],0)
        self.assertEqual(second['channels']['market_demand']['accepted'],0)
        self.assertEqual(intake.cycle(self.root,supply=supply)['enqueued'],0)
    def test_mathematical_owner_briefs_have_collision_free_source_ids(self):
        config={'schema':'saeit.factory.internal-design-supply.v1','briefs':[
            {'schema':intake.SCHEMA,'family_id':'F28','mathematical_definition':{
                'title':'Measured inventory gap','expression':'needed - held',
                'inputs':[{'name':'needed','unit':'units','minimum':0,'maximum':100},
                          {'name':'held','unit':'units','minimum':0,'maximum':100}],
                'output_unit':'units'}},
            {'schema':intake.SCHEMA,'family_id':'F28','mathematical_definition':{
                'title':'Measured liquidity gap','expression':'required - cash',
                'inputs':[{'name':'required','unit':'currency','minimum':0,'maximum':100},
                          {'name':'cash','unit':'currency','minimum':0,'maximum':100}],
                'output_unit':'currency'}}]}
        source=self.root/'internal.json'
        designer.save(source,config)
        result=intake.cycle(self.root,supply=source)
        self.assertEqual(result['enqueued'],2)
        ids=[row['source_id'] for row in result['latest_results']]
        self.assertEqual(len(set(ids)),2)
        self.assertEqual(intake.cycle(self.root,supply=source)['enqueued'],0)

    def test_record_cannot_spoof_internal_exemption(self):
        r=self.record('site_need');r['internal_math_prototype']=True
        self.snapshot('site-topic-snapshot',[r])
        result=intake.cycle(self.root)
        self.assertEqual(result['enqueued'],0)
        self.assertIn('INTERNAL_SOURCE_SPOOF',result['latest_results'][0]['reason'])
    def test_site_bridge_receipt_is_checked_before_intake(self):
        p=self.root/'production-instructions/orders/WO-00000001.json'
        r=self.record('site_need');r.update(id='WO-00000001',origin='site-topic-snapshot',requesting_site='zomorodmelal.ir',bridge_cycle='2026-10-19')
        designer.save(p,r)
        result=intake.cycle(self.root)
        self.assertEqual(result['enqueued'],0)
        designer.save(self.root/'production-instructions/zomorod-site-order-receipts/2026-10-19.json',{'schema':'saeit.zomorod.site-factory-bridge.v1','site':'zomorodmelal.ir','orders':['WO-00000001']})
        result=intake.cycle(self.root,now=10**12)
        self.assertEqual(result['enqueued'],1)

    def test_research_brief_does_not_bypass_the_existing_evidence_gate(self):
        path=self.root/'research-dossiers/product-1/evidence_manifest.json'
        designer.save(path,{'product_id':1,'engineering_brief':{'schema':intake.SCHEMA,'blueprint':'inventory_coverage'},
                            'sources':[],'evidence':[],'reports':[],'findings':[],'issues':['missing evidence'],'missing_factors':[]})
        result=intake.cycle(self.root)
        self.assertEqual(result['enqueued'],0)
        self.assertIn('RESEARCH_EVIDENCE_NOT_READY',result['latest_results'][0]['reason'])

    def test_provenance_failure_does_not_overwrite_previous_verified_product(self):
        spec=intake.expand({'schema':intake.SCHEMA,'blueprint':'inventory_coverage'},'factory_invention')
        result=designer.compile_recipe(spec,self.root/'recipes');recipe=result['recipe']
        output=self.root/'digital-products';maker.manufacture(recipe,self.root/'recipes',output)
        before=(output/recipe/(recipe+'-internal-qa.zip')).read_bytes()
        path=self.root/'recipes'/recipe/'design_receipt.json';receipt=json.loads(path.read_text());receipt['request_sha256']='a'*64
        path.write_text(json.dumps(receipt))
        with self.assertRaisesRegex(RuntimeError,'PROVENANCE_MISMATCH'):maker.manufacture(recipe,self.root/'recipes',output)
        self.assertEqual((output/recipe/(recipe+'-internal-qa.zip')).read_bytes(),before)

    def test_new_arithmetic_requirement_needs_no_preexisting_blueprint_or_recipe(self):
        record=self.record('site_need')
        record['design_brief']={'schema':intake.SCHEMA,'family_id':'F28','mathematical_definition':{
            'title':'Replenished stock coverage','expression':'(stock + incoming) / daily_use',
            'inputs':[{'name':'stock','unit':'units','minimum':0,'maximum':100},
                      {'name':'incoming','unit':'units','minimum':0,'maximum':100},
                      {'name':'daily_use','unit':'units/day','minimum':1,'maximum':100}],
            'output_unit':'days'}}
        self.snapshot('site-topic-snapshot',[record])
        self.assertEqual(intake.cycle(self.root)['enqueued'],1)
        result=designer.cycle(self.root)
        self.assertEqual(result['new_recipes'],1)
        recipe=result['actions'][0]['result']['recipe']
        source=self.root/'generated-product-recipes'/recipe
        self.assertEqual(maker.verify_recipe(source),'unit_tests_passed')
        spec=importlib.util.spec_from_file_location('new_arithmetic',source/'calculator.py')
        program=importlib.util.module_from_spec(spec);spec.loader.exec_module(program)
        self.assertEqual(program.calculate({'stock':80,'incoming':40,'daily_use':10})['result'],'12')

    def test_declared_expression_cannot_execute_python_or_use_unsupported_operators(self):
        base={'schema':intake.SCHEMA,'mathematical_definition':{
            'inputs':[{'name':'amount','unit':'units'}],'output_unit':'units'}}
        for expression in ("__import__('os').getcwd()",'amount**amount','amount[0]','1/amount'):
            base['mathematical_definition']['expression']=expression
            with self.assertRaises(ValueError):intake.expand(base,'factory_invention')
