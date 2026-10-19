#!/usr/bin/env python3
"""Real private artifact checks; installation is separate from product quality."""
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import subprocess
import sys
import time
import tempfile


def run(command, *, timeout=90, env=None):
    result=subprocess.run([str(x) for x in command],capture_output=True,text=True,
                          timeout=timeout,env=env)
    if result.returncode:
        details=result.stderr+result.stdout
        raise RuntimeError('COMMAND_FAILED:'+str(command[0])+':'+details[:700]+'\n'+details[-800:])
    return result.stdout


def prove(root):
    root=Path(root); work=Path(tempfile.mkdtemp(prefix='proof-',dir=str(root)))
    checks={}; artifacts={}
    def check(name, callback):
        started=time.monotonic()
        try:
            evidence=callback()
            checks[name]={'passed':True,'evidence':evidence,'seconds':round(time.monotonic()-started,3)}
        except Exception as error:
            checks[name]={'passed':False,'error':str(error)[-1800:]}
    def artifact(path):
        if not path.is_file() or path.stat().st_size < 20:
            raise ValueError('MISSING_REAL_ARTIFACT:'+str(path))
        artifacts[path.name]={'bytes':path.stat().st_size,
                              'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
        return artifacts[path.name]
    def images():
        from PIL import Image
        file=work/'image.png'; Image.new('RGB',(128,128),(12,120,40)).save(file)
        with Image.open(file) as image:
            image.verify()
        return artifact(file)
    def documents():
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches
        from openpyxl import Workbook,load_workbook
        from reportlab.pdfgen import canvas
        from pypdf import PdfReader
        doc=Document();doc.add_heading('Factory installation verification',0)
        doc.add_paragraph('This is a technical test artifact, not a commercial product.')
        doc.save(work/'document.docx')
        assert 'technical test' in Document(work/'document.docx').paragraphs[1].text
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text_frame.text='Factory verification'
        deck.save(work/'slides.pptx');assert len(Presentation(work/'slides.pptx').slides)==1
        workbook=Workbook();workbook.active['A1']=2;workbook.active['A2']=3;workbook.active['A3']='=SUM(A1:A2)'
        workbook.save(work/'sheet.xlsx');assert load_workbook(work/'sheet.xlsx').active['A3'].value=='=SUM(A1:A2)'
        pdf=canvas.Canvas(str(work/'document.pdf'));pdf.drawString(60,780,'Factory verification');pdf.save()
        assert 'Factory verification' in PdfReader(work/'document.pdf').pages[0].extract_text()
        return {p.name:artifact(p) for p in (work/'document.docx',work/'slides.pptx',work/'sheet.xlsx',work/'document.pdf')}
    def media():
        target=work/'video.mp4'
        run([root/'bin/ffmpeg','-hide_banner','-loglevel','error','-y','-f','lavfi','-i',
             'testsrc=size=128x128:rate=8','-t','1','-threads','1','-pix_fmt','yuv420p',target])
        probe=json.loads(run([root/'bin/ffprobe','-v','error','-show_streams','-of','json',target]))
        assert probe['streams'][0]['width']==128
        return artifact(target)
    def office():
        from pypdf import PdfReader
        converted=work/'converted';converted.mkdir(exist_ok=True)
        for name in ('document.docx','slides.pptx','sheet.xlsx'):
            run([root/'bin/libreoffice','-env:UserInstallation='+str((work/'lo-profile').as_uri()),
                 '--headless','--convert-to','pdf','--outdir',converted,work/name],timeout=100)
            result=converted/(Path(name).stem+'.pdf')
            assert len(PdfReader(result).pages)>=1
            artifact(result)
        # Recalculate an actual spreadsheet formula using LibreOffice.
        run([root/'bin/libreoffice','-env:UserInstallation='+str((work/'lo-profile').as_uri()),
             '--headless','--convert-to','xlsx','--outdir',converted,work/'sheet.xlsx'],timeout=90)
        from openpyxl import load_workbook
        assert load_workbook(converted/'sheet.xlsx',data_only=True).active['A3'].value==5
        return {'pdf_conversions':3,'spreadsheet_formula_result':5}
    def render():
        from PIL import Image
        script=work/'render.py'
        script.write_text("import bpy\ns=bpy.context.scene\ns.render.engine='CYCLES'\ns.cycles.samples=1\ns.render.threads_mode='FIXED'\ns.render.threads=1\ns.render.resolution_x=64\ns.render.resolution_y=64\ns.render.resolution_percentage=100\ns.render.filepath="+repr(str(work/'cube.png'))+"\nbpy.ops.wm.save_as_mainfile(filepath="+repr(str(work/'cube.blend'))+")\nbpy.ops.render.render(write_still=True)\n")
        run([root/'bin/blender','-b','--factory-startup','--python',script],timeout=180)
        with Image.open(work/'cube.png') as image: assert image.size==(64,64)
        artifact(work/'cube.blend')
        return artifact(work/'cube.png')
    def llm():
        import socket, urllib.request, urllib.error
        with socket.socket() as reserved:
            reserved.bind(('127.0.0.1',0)); port=reserved.getsockname()[1]
        logfile=(work/'local-model-server.log').open('w')
        command=[str(root/'llama/bin/llama-server'),'-m',str(root/'models/coder.gguf'),
                 '--host','127.0.0.1','--port',str(port),'-t','1','-c','512',
                 '--parallel','1','--no-warmup']
        process=None
        try:
            process=subprocess.Popen(command,stdout=logfile,stderr=logfile)
            deadline=time.monotonic()+100
            while time.monotonic()<deadline:
                if process.poll() is not None: raise ValueError('LOCAL_MODEL_SERVER_EXITED')
                try:
                    with urllib.request.urlopen('http://127.0.0.1:'+str(port)+'/health',timeout=2) as response:
                        if response.status==200: break
                except (urllib.error.URLError,TimeoutError): pass
                time.sleep(0.25)
            else: raise ValueError('LOCAL_MODEL_START_TIMEOUT')
            payload={'messages':[{'role':'system','content':'Reply to arithmetic with only the numerical result.'},
                                 {'role':'user','content':'2+2'}],
                     'max_tokens':16,'temperature':0,'seed':42,'stream':False}
            request=urllib.request.Request('http://127.0.0.1:'+str(port)+'/v1/chat/completions',
                data=json.dumps(payload).encode(),headers={'Content-Type':'application/json'})
            with urllib.request.urlopen(request,timeout=100) as response:
                result=json.load(response)
            output=result['choices'][0]['message']['content'].strip()
            (work/'local-model-output.txt').write_text(output)
            if output!='4': raise ValueError('LOCAL_MODEL_REFERENCE_PROMPT_FAILED:'+output[-200:])
            return {'reference_answer':output,'output_sha256':hashlib.sha256(output.encode()).hexdigest(),
                    'completion_tokens':result.get('usage',{}).get('completion_tokens'),
                    'max_generated_tokens':16,'loopback_only':True,
                    'quality_of_other_answers_verified':False}
        finally:
            if process is not None and process.poll() is None:
                process.terminate()
                try: process.wait(timeout=10)
                except subprocess.TimeoutExpired: process.kill(); process.wait(timeout=5)
            logfile.close()
    for name,callback in [('images',images),('documents',documents),('video',media),
                          ('office_render_and_recalculation',office),('blender_cpu_render',render),('local_language_model',llm)]:
        check(name,callback)
    report={'checks':checks,'artifacts':artifacts,
            'packages':{name:importlib.metadata.version(name) for name in ('Pillow','python-docx','python-pptx','reportlab','openpyxl','pypdf')},
            'product_quality_score':None,'generation_promotions':0}
    (root/'proof.json').write_text(json.dumps(report,indent=2,sort_keys=True)+'\n')
    print(json.dumps({name:row['passed'] for name,row in checks.items()},sort_keys=True))
    if not all(row['passed'] for row in checks.values()): raise SystemExit(1)
    return report

if __name__=='__main__':prove(sys.argv[1])
