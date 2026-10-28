#!/usr/bin/env python3
"""Source-locked WDI 2024->2025 *algebraic* component-revision audit.

The agricultural value-added-per-worker series P and constant-2015-dollar
agricultural value added V imply a denominator L*=V/P. L* is NOT an
independently validated ILO employment series. For each country,
delta growth(P) = delta growth(V) - delta growth(L*) is an identity, not a
causal explanation. Original World Bank archives, never fabricated values.
"""
import csv
import hashlib
import importlib.util
import json
import math
from pathlib import Path
import re
import statistics
import tempfile
import zipfile

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
PANEL = ROOT / "AE-2026-001_OFFICIAL_WDI_OBSERVED_COUNTRY_PANEL_2026-10-25.csv"
OUTPUT = ROOT / "AE-2026-001_2024_2025_COMPONENT_REVISION_IDENTITY_AUDIT_2026-10-26.json"
SOURCE = HERE / "official_wdi_2024_vs_2026_vintage_audit.py"
spec = importlib.util.spec_from_file_location("official_vintage", SOURCE)
base = importlib.util.module_from_spec(spec)
spec.loader.exec_module(base)

ARCHIVES = {
    "2024": ("https://databankfiles.worldbank.org/public/ddpext_download/Archive/WDI_excel_2024_10_24.zip",
             "5a8922c6629eea69b93a314ebb2363ed8d5faf4f345ff19dbe203c62692af5a1"),
    "2025": ("https://databankfiles.worldbank.org/public/ddpext_download/Archive/WDI_excel_2025_10_08.zip",
             "9185df6f3029ca6c990d87674dcec60e4633cf0b344bc4ee6459392f1cfdc173"),
}
INDICATORS = {"NV.AGR.EMPL.KD": "per_worker", "NV.AGR.TOTL.KD": "value_added"}
YEARS = (2000, 2019)

def read_archive(path, scratch):
    result = {name: {} for name in INDICATORS.values()}
    scanned = []
    with zipfile.ZipFile(path) as z:
        candidates = [x for x in z.infolist() if x.filename.lower().endswith(".xlsx")
                      and not x.is_dir() and x.file_size <= 1_000_000_000]
        if not candidates:
            raise RuntimeError("Official archive has no eligible XLSX")
        for index, item in enumerate(candidates):
            local = scratch / f"extract_{index}.xlsx"
            with z.open(item) as src, local.open("wb") as dst:
                import shutil
                shutil.copyfileobj(src, dst, length=4*1024*1024)
            wb = load_workbook(local, read_only=True, data_only=True)
            try:
                for sheet in wb.worksheets:
                    header = None
                    found_in_sheet = 0
                    for row in sheet.iter_rows(values_only=True):
                        cells = list(row)
                        if header is None:
                            first = [str(v or "").strip().lower() for v in cells[:6]]
                            if "country code" in first and "indicator code" in first:
                                ycols = {year: next((j for j, v in enumerate(cells)
                                        if base.year_from_cell(v) == year), None) for year in YEARS}
                                if all(j is not None for j in ycols.values()):
                                    header = (first.index("country code"),
                                              first.index("indicator code"),
                                              first.index("indicator name") if "indicator name" in first else None,
                                              ycols)
                            continue
                        co, ind, label, ycols = header
                        if max(co, ind, *ycols.values()) >= len(cells):
                            continue
                        indicator = str(cells[ind] or "").strip()
                        if indicator not in INDICATORS:
                            continue
                        iso = str(cells[co] or "").strip()
                        if not re.fullmatch(r"[A-Z0-9]{3}", iso):
                            continue
                        if label is None or label >= len(cells):
                            raise RuntimeError("Indicator label missing; cannot verify base")
                        name = str(cells[label] or "").lower()
                        if "2015" not in name or "agriculture" not in name:
                            raise RuntimeError(f"Unexpected indicator definition {indicator}: {name}")
                        if indicator == "NV.AGR.EMPL.KD" and "worker" not in name:
                            raise RuntimeError("Per-worker label mismatch")
                        if indicator == "NV.AGR.TOTL.KD" and "worker" in name:
                            raise RuntimeError("Value-added numerator label mismatch")
                        values = {str(y): base.maybe_pos(cells[ycols[y]]) for y in YEARS}
                        target = result[INDICATORS[indicator]]
                        if iso in target and target[iso] != values:
                            raise RuntimeError(f"Conflicting duplicate {indicator} {iso}")
                        target[iso] = values
                        found_in_sheet += 1
                    scanned.append({"xlsx": item.filename, "sheet": sheet.title,
                                    "target_rows": found_in_sheet})
            finally:
                wb.close()
                local.unlink(missing_ok=True)
    if any(len(v) < 120 for v in result.values()):
        raise RuntimeError("Too few historical rows for either source indicator")
    return result, scanned

def rate(row):
    if not row:
        return None
    a, b = row.get("2000"), row.get("2019")
    return math.log(b / a) / 19 if a is not None and b is not None else None

def summary(values):
    return {"n": len(values), "median": statistics.median(values) if values else None,
            "median_abs": statistics.median(abs(x) for x in values) if values else None}

def main():
    current = {}
    with PANEL.open(newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            iso = row["country_iso3"]
            current[iso] = {"2000": base.maybe_pos(row[base.COLUMN[2000]]),
                            "2019": base.maybe_pos(row[base.COLUMN[2019]])}
    if len(current) != 217:
        raise RuntimeError("Frozen 2026 country snapshot must contain 217 economy codes")

    datasets, provenance = {}, {}
    with tempfile.TemporaryDirectory(prefix="wdi_component_") as folder:
        scratch = Path(folder)
        for vintage, (url, expected_sha) in ARCHIVES.items():
            base.ARCHIVE = url
            destination = scratch / f"{vintage}.zip"
            meta = base.fetch_archive(destination)
            if meta["sha256"] != expected_sha:
                raise RuntimeError(f"Historical {vintage} SHA mismatch: fail closed")
            datasets[vintage], scanned = read_archive(destination, scratch)
            provenance[vintage] = {**meta, "sheets": scanned}
            destination.unlink()

    cohort = sorted(set(current) & set(datasets["2024"]["per_worker"])
                    & set(datasets["2025"]["per_worker"]))
    rows = []
    for iso in cohort:
        if rate(current[iso]) is None:
            continue
        pair = {}
        for vintage in ("2024", "2025"):
            per = datasets[vintage]["per_worker"].get(iso)
            val = datasets[vintage]["value_added"].get(iso)
            gp, gv = rate(per), rate(val)
            if gp is None or gv is None:
                pair = {}
                break
            implied = {str(y): val[str(y)] / per[str(y)] for y in YEARS}
            gl = rate(implied)
            if abs(gp - (gv - gl)) > 1e-11:
                raise RuntimeError("Within-vintage logarithmic identity failed")
            pair[vintage] = {"per_worker_growth": gp, "value_added_growth": gv,
                             "implied_denominator_growth": gl}
        if len(pair) != 2:
            continue
        dp = pair["2025"]["per_worker_growth"] - pair["2024"]["per_worker_growth"]
        dv = pair["2025"]["value_added_growth"] - pair["2024"]["value_added_growth"]
        dl = pair["2025"]["implied_denominator_growth"] - pair["2024"]["implied_denominator_growth"]
        if abs(dp - (dv - dl)) > 1e-11:
            raise RuntimeError("Cross-vintage revision identity failed")
        rows.append({"iso3": iso, "2024": pair["2024"], "2025": pair["2025"],
                     "delta_growth_per_worker": dp, "delta_growth_value_added": dv,
                     "delta_growth_implied_denominator": dl})

    if len(rows) < 100:
        raise RuntimeError(f"Only {len(rows)} jointly observed source component cases")
    dp = [r["delta_growth_per_worker"] for r in rows]
    dv = [r["delta_growth_value_added"] for r in rows]
    dl = [r["delta_growth_implied_denominator"] for r in rows]
    report = {
        "question": "How do revisions in the published value-added series and its implied residual denominator algebraically partition per-worker growth revisions?",
        "indicator_units": "Both WDI indicator labels must say constant 2015 US$; P is per worker; V is total agricultural value added.",
        "archive_provenance": provenance,
        "frozen_2026_panel_sha256": hashlib.sha256(PANEL.read_bytes()).hexdigest(),
        "year_pair": [2000, 2019], "country_count": len(rows),
        "country_level_observations": rows,
        "descriptive_revision_statistics": {
            "per_worker": summary(dp), "value_added": summary(dv),
            "implied_denominator": summary(dl),
            "value_added_abs_revision_larger_than_implied_denominator_count":
                sum(abs(r["delta_growth_value_added"]) >
                    abs(r["delta_growth_implied_denominator"]) for r in rows)
        },
        "quality_gates": {
            "archive_sha256": "PASS", "base_year_labels": "PASS",
            "identity_residual_max": max(abs(r["delta_growth_per_worker"] -
                (r["delta_growth_value_added"] - r["delta_growth_implied_denominator"]))
                for r in rows),
            "publication_novelty": "NOT_ESTABLISHED",
            "causal_source_attribution": "NOT_IDENTIFIED"
        },
        "limitations": [
            "L*=V/P is an algebraically implied denominator, NOT a separately observed or independently verified ILO employment series.",
            "The identity is exact by construction and does not identify causal sources of revision.",
            "The joint sample may be smaller than the 155-country three-vintage per-worker cohort; never silently equate them.",
            "No 2026 component series are fetched here; the audit compares original 2024 and 2025 archives only.",
            "Changes in national-account estimates, exchange-rate conversion, ILO modeled employment or definitions cannot be causally apportioned by this exercise.",
            "No confidence interval, significance test, causal effect, policy inference or publication novelty is claimed."
        ]
    }
    OUTPUT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"country_count": len(rows),
        "descriptive_revision_statistics": report["descriptive_revision_statistics"],
        "quality_gates": report["quality_gates"]}, indent=2))

if __name__ == "__main__":
    main()
