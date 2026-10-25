#!/usr/bin/env python3
"""Real WDI historical-vintage audit: October 2024 XLSX vs October 2026 WDI.

Direct original-source extract, not a significance test and not proof of novelty.
World Bank archive is explicitly linked on the World Development Indicators
WDI Archives page. Validate indicator definition and price base before comparing.
No missing values are filled with zero or borrowed from another vintage.
"""
import csv
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import statistics
import tempfile
import urllib.request
import zipfile

ARCHIVE = "https://databankfiles.worldbank.org/public/ddpext_download/Archive/WDI_excel_2024_10_24.zip"
INDICATOR = "NV.AGR.EMPL.KD"
YEARS = (2000, 2019, 2023)
ROOT = Path(__file__).resolve().parent.parent
CURRENT = ROOT / "AE-2026-001_OFFICIAL_WDI_OBSERVED_COUNTRY_PANEL_2026-10-25.csv"
COLUMN = {
    2000: "ag_value_added_per_worker_2000_constant2015usd",
    2019: "ag_value_added_per_worker_2019_constant2015usd",
    2023: "ag_value_added_per_worker_2023_constant2015usd",
}
OUTPUT = ROOT / "AE-2026-001_OFFICIAL_2024_VS_2026_WDI_VINTAGE_EMPIRICAL_AUDIT_2026-10-26.json"

def maybe_pos(x):
    try:
        z = float(x)
    except (TypeError, ValueError):
        return None
    return z if math.isfinite(z) and z > 0 else None

def fetch_archive(dest):
    req = urllib.request.Request(ARCHIVE, headers={"User-Agent":"Academic-WDI-Vintage-Reproducibility/1.0"})
    h = hashlib.sha256()
    total=0
    with urllib.request.urlopen(req, timeout=120) as src, dest.open("wb") as f:
        while chunk := src.read(4*1024*1024):
            total += len(chunk)
            if total > 110_000_000:
                raise RuntimeError("Archive exceeds 110 MB safety cap")
            f.write(chunk);h.update(chunk)
    if total < 5_000_000:
        raise RuntimeError("Archive response implausibly small")
    if not zipfile.is_zipfile(dest):
        raise RuntimeError("Official URL response not a valid ZIP")
    return {"url":ARCHIVE, "size_bytes":total, "sha256":h.hexdigest()}

def year_from_cell(x):
    s=str(x or "").strip()
    m=re.match(r"^(19|20)\d{2}(?:\s*\[YR\d{4}\])?$",s)
    return int(s[:4]) if m else None

def load_original_wdi_excel(zpath, scratch):
    from openpyxl import load_workbook
    found={}
    checked=[]
    with zipfile.ZipFile(zpath) as archive:
        candidates=[m for m in archive.infolist() if
                    m.filename.lower().endswith(".xlsx") and not m.is_dir()
                    and m.file_size <= 1_000_000_000]
        if not candidates:
            raise RuntimeError("No safely sized Excel .xlsx in official archive; inspect archive format")
        for ix,item in enumerate(candidates):
            local=scratch / ("wdi_"+str(ix)+".xlsx")
            with archive.open(item) as src, local.open("wb") as out:
                import shutil
                shutil.copyfileobj(src,out,length=4*1024*1024)
            wb=load_workbook(local,read_only=True,data_only=True)
            try:
                for sheet in wb.worksheets:
                    header=None
                    for row in sheet.iter_rows(values_only=True):
                        values=list(row)
                        if header is None:
                            txt=[str(q or "").strip().lower() for q in values[:6]]
                            if ("country code" in txt and "indicator code" in txt):
                                yrs={y:next((i for i,v in enumerate(values) if year_from_cell(v)==y),None)
                                     for y in YEARS}
                                if all(i is not None for i in yrs.values()):
                                    header=(txt.index("country code"),txt.index("indicator code"),yrs,
                                            txt.index("indicator name") if "indicator name" in txt else None)
                            continue
                        co,ind,yrs,namecol=header
                        if max([co,ind,*yrs.values()])>=len(values):
                            continue
                        if str(values[ind] or "").strip()!=INDICATOR:
                            continue
                        iso=str(values[co] or "").strip()
                        if not re.fullmatch(r"[A-Z0-9]{3}",iso):
                            continue
                        indicator_name=str(values[namecol] or "") if namecol is not None and namecol<len(values) else ""
                        if "2015" not in indicator_name or "worker" not in indicator_name.lower():
                            raise RuntimeError("Indicator label not explicitly constant-2015 USD per worker: "+indicator_name)
                        vals={str(y):maybe_pos(values[yrs[y]]) for y in YEARS}
                        if iso in found and found[iso]!=vals:
                            raise RuntimeError("Nonmatching duplicated old-vintage country "+iso)
                        found[iso]=vals
                    checked.append({"xlsx":item.filename,"sheet":sheet.title,"country_rows_found":len(found)})
            finally:
                wb.close();local.unlink(missing_ok=True)
    if len(found)<120:
        raise RuntimeError(f"Only {len(found)} historical indicator rows; refuse incomplete vintage comparison")
    return found, checked

def compare(old):
    current={}
    with CURRENT.open(newline="",encoding="utf-8") as f:
        for row in csv.DictReader(f):
            iso=row["country_iso3"]
            current[iso]={str(y):maybe_pos(row[COLUMN[y]]) for y in YEARS}
    if len(current)!=217:
        raise RuntimeError("2026 archived source country count changed unexpectedly")
    aligned=sorted(set(old)&set(current))
    result={"country_set_2026":len(current),"historical_indicator_rows":len(old),
            "overlapping_codes":len(aligned),"by_year":{},"by_growth_period":{}}
    if len(aligned)<150:
        raise RuntimeError("Unexpectedly few overlapping ISO3 codes")
    for y in YEARS:
        pairs=[(iso,old[iso][str(y)],current[iso][str(y)])
               for iso in aligned if old[iso][str(y)] is not None and current[iso][str(y)] is not None]
        logs=[math.log(z/new) for _,z,new in pairs]
        result["by_year"][str(y)]={
            "matched_nonmissing":len(pairs),
            "median_log_old_over_new":statistics.median(logs) if logs else None,
            "median_absolute_log_revision":statistics.median(abs(v) for v in logs) if logs else None,
            "changed_more_than_1pct_log":sum(abs(v)>0.01 for v in logs),
            "source_coverage_old":sum(old[iso][str(y)] is not None for iso in aligned),
            "source_coverage_2026":sum(current[iso][str(y)] is not None for iso in aligned)}
    for end in (2019,2023):
        old_growth={iso:math.log(old[iso][str(end)]/old[iso]["2000"])/(end-2000)
                    for iso in aligned if old[iso]["2000"] is not None and old[iso][str(end)] is not None}
        new_growth={iso:math.log(current[iso][str(end)]/current[iso]["2000"])/(end-2000)
                    for iso in aligned if current[iso]["2000"] is not None and current[iso][str(end)] is not None}
        both=sorted(set(old_growth)&set(new_growth))
        if len(both)<100:
            raise RuntimeError("Insufficient common vintage growth pairs")
        d=[new_growth[iso]-old_growth[iso] for iso in both]
        result["by_growth_period"]["2000_"+str(end)]={
            "old_vintage_available_growth_pairs":len(old_growth),
            "new_vintage_available_growth_pairs":len(new_growth),
            "common_country_pairs":len(both),
            "old_median_growth_common":statistics.median(old_growth[iso] for iso in both),
            "new_median_growth_common":statistics.median(new_growth[iso] for iso in both),
            "median_individual_annual_growth_revision":statistics.median(d),
            "sign_change_of_annual_growth":sum((old_growth[iso]>0)!=(new_growth[iso]>0) for iso in both)}
    return result

def main():
    with tempfile.TemporaryDirectory(prefix="wdi_vintage_") as tmp:
        scratch=Path(tmp)
        meta=fetch_archive(scratch/"official_archive.zip")
        older,sheets=load_original_wdi_excel(scratch/"official_archive.zip",scratch)
    report={
      "comparison":"Official World Bank 2024-10-24 historical Excel archive versus source-archived official WDI 2026-10-25 API",
      "retrieved_utc":datetime.now(timezone.utc).isoformat(),
      "indicator":INDICATOR,
      "units":"agriculture, forestry and fishing value added per worker; constant 2015 USD",
      "archive":meta,
      "archive_sheets_inspected":sheets,
      "outcome":compare(older),
      "limitations":["World Bank warns historical WDI estimates are revised; differing country-years must not be backfilled.",
                     "Only source-confirmed same-definition/same-price-base indicators may be compared.",
                     "Revisions can reflect ILO-modelled employment and national accounts, not measured farm productivity changes.",
                     "Methodological evidence of vintage sensitivity does not itself prove first-ever or publication novelty."]
    }
    OUTPUT.write_text(json.dumps(report,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps(report,indent=2,ensure_ascii=False))

if __name__=="__main__":main()
