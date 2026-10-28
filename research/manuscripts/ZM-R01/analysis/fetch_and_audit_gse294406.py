#!/usr/bin/env python3
"""Acquire original GEO GSE294406 processed FPKM workbook, validate workbook, record SHA256.
This does NOT create raw gene counts and does NOT perform differential-expression analysis.
Source: https://ftp.ncbi.nlm.nih.gov/geo/series/GSE294nnn/GSE294406/suppl/
"""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import time
import urllib.request
import zipfile

URL = "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE294nnn/GSE294406/suppl/GSE294406_FPKM_all_treatments.xlsx"
NAME = "GSE294406_FPKM_all_treatments.xlsx"
LIMIT = 100 * 1024 * 1024
REQUIRED = {"study", "accession", "condition", "drought", "heat", "co2", "biological_replicate"}


def validate_manifest(path):
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if len(rows) != 24 or not REQUIRED.issubset(rows[0]):
        raise ValueError("GSE294406 manifest must contain exactly 24 samples and required columns")
    seen = set()
    groups = {}
    for row in rows:
        gsm = row["accession"]
        if gsm in seen or row["study"] != "GSE294406":
            raise ValueError("Duplicate accession or wrong study")
        seen.add(gsm)
        key = (int(row["drought"]), int(row["heat"]), int(row["co2"]))
        groups.setdefault(key, set()).add(int(row["biological_replicate"]))
    if len(groups) != 8 or any(v != {1, 2, 3} for v in groups.values()):
        raise ValueError("2x2x2 study requires three independent biological samples per condition")
    return {"samples": len(rows), "factorial_cells": len(groups),
            "biological_replicates_per_cell": 3, "factorial_parameters": 8,
            "full_design_residual_df": 16, "GSM_ids": sorted(seen)}


def download(path, attempts=3):
    if path.exists() and path.stat().st_size > 0:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    for attempt in range(attempts):
        try:
            request = urllib.request.Request(URL, headers={"User-Agent": "ZM-R01-reproducible-research/1.0 (public GEO accession)"})
            with urllib.request.urlopen(request, timeout=75) as response, path.with_suffix(".part").open("wb") as out:
                n = 0
                while True:
                    block = response.read(1024 * 1024)
                    if not block:
                        break
                    n += len(block)
                    if n > LIMIT:
                        raise ValueError("File exceeds 100 MiB safety limit")
                    out.write(block)
            path.with_suffix(".part").replace(path)
            return
        except Exception:
            path.with_suffix(".part").unlink(missing_ok=True)
            if attempt == attempts - 1:
                raise
            time.sleep(3 * (attempt + 1))


def audit_workbook(path):
    if not zipfile.is_zipfile(path):
        raise ValueError("Downloaded file is not a valid XLSX ZIP container")
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet in book:
        preview = []
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 0, 4),
                                   max_col=min(sheet.max_column or 0, 30), values_only=True):
            preview.append([str(x)[:100] if x is not None else None for x in row])
        sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                       "top_left_preview": preview})
    book.close()
    return sheets


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    parser.add_argument("--manifest-only", action="store_true",
                        help="Validate manifest without downloading or claiming expression analysis")
    args = parser.parse_args()
    meta = validate_manifest(args.manifest)
    args.outdir.mkdir(parents=True, exist_ok=True)
    report = {"accession": "GSE294406", "manifest": meta,
              "source_url": URL, "source_format": "FPKM processed XLSX",
              "important_limitation": "FPKM values are normalized expression, NOT raw integer counts. Do NOT pass FPKM to DESeq2/edgeR count-based differential-expression models. Obtain SRA FASTQ or a genuine raw count matrix and quantify consistently.",
              "comparison_warning": "Use within-study factorial contrasts; do NOT pool with GSE157033 or GSE226842 as replicates.",
              "analysis_performed": "manifest metadata validation only"}
    if not args.manifest_only:
        workbook = args.outdir / NAME
        download(workbook)
        h = hashlib.sha256()
        with workbook.open("rb") as fh:
            for chunk in iter(lambda: fh.read(1024 * 1024), b""):
                h.update(chunk)
        report.update({"file_size_bytes": workbook.stat().st_size,
                       "sha256": h.hexdigest(),
                       "workbook_sheets": audit_workbook(workbook),
                       "analysis_performed": "actual GEO workbook download, checksum and spreadsheet-structure inspection; NO raw counts, DEGs or mechanism inferred"})
    target = args.outdir / "GSE294406_acquisition_audit.json"
    target.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"audit_path": str(target), "analysis_performed": report["analysis_performed"],
                      "sha256": report.get("sha256"), "manifest": meta}, indent=2))


if __name__ == "__main__":
    main()
